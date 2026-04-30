"""Compute funnel metrics from career-scan inputs."""
from __future__ import annotations

import json
import logging
import statistics
from collections import Counter
from datetime import date, datetime, timedelta, timezone
from pathlib import Path

import openpyxl

logger = logging.getLogger(__name__)

SHEET_NAME = "Remote Jobs"
TERMINAL_STATES = {"submitted", "responded", "rejected", "skipped"}


def _today() -> date:
    return datetime.now(timezone.utc).date()


def _parse_iso(value) -> date | None:
    if isinstance(value, date) and not isinstance(value, datetime):
        return value
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, str) and value:
        try:
            return date.fromisoformat(value[:10])
        except ValueError:
            return None
    return None


def _within_days(d: date | None, days: int) -> bool:
    if d is None:
        return False
    return (_today() - d).days < days


def scrape_metrics(excel_path: Path) -> dict:
    wb = openpyxl.load_workbook(excel_path, read_only=True)
    if SHEET_NAME not in wb.sheetnames:
        return {"error": f"sheet {SHEET_NAME!r} missing"}
    ws = wb[SHEET_NAME]
    rows = list(ws.iter_rows(min_row=1, values_only=True))
    if len(rows) < 2:
        return {"total": 0, "added_24h": 0, "added_7d": 0, "by_source": {}, "top_companies": []}

    header = [str(c) if c is not None else "" for c in rows[0]]
    idx = {name: i for i, name in enumerate(header)}

    total = len(rows) - 1
    added_24h = added_7d = 0
    by_source: Counter = Counter()
    by_company: Counter = Counter()

    for row in rows[1:]:
        added = _parse_iso(row[idx.get("Date Added", -1)] if idx.get("Date Added") is not None else None)
        source = (row[idx["Source"]] or "") if "Source" in idx else ""
        company = (row[idx["Company"]] or "") if "Company" in idx else ""
        if _within_days(added, 1):
            added_24h += 1
        if _within_days(added, 7):
            added_7d += 1
        by_source[str(source)] += 1
        if company:
            by_company[str(company)] += 1

    return {
        "total": total,
        "added_24h": added_24h,
        "added_7d": added_7d,
        "by_source": dict(by_source.most_common()),
        "top_companies": by_company.most_common(15),
    }


def _ts(s: str) -> datetime | None:
    try:
        return datetime.fromisoformat(s)
    except (TypeError, ValueError):
        return None


def _hours_between(a: str, b: str) -> float | None:
    da, db = _ts(a), _ts(b)
    if da is None or db is None:
        return None
    return (db - da).total_seconds() / 3600.0


def application_metrics(state_path: Path) -> dict:
    if not state_path.exists() or state_path.stat().st_size == 0:
        return {"total": 0, "by_state": {}, "errors": [], "median_tailor_to_submit_hours": None}

    state = json.loads(state_path.read_text())
    by_state: Counter = Counter()
    error_reasons: Counter = Counter()
    tailor_to_submit: list[float] = []
    submit_to_response: list[float] = []
    submitted_24h = submitted_7d = 0

    for rec in state.values():
        by_state[rec["state"]] += 1
        if rec.get("error"):
            error_reasons[rec["error"][:140]] += 1
        history = rec.get("history") or []
        # find first transitions into each state
        first_at: dict[str, str] = {}
        for h in history:
            first_at.setdefault(h["to"], h["at"])
        if "tailored" in first_at and "submitted" in first_at:
            hours = _hours_between(first_at["tailored"], first_at["submitted"])
            if hours is not None:
                tailor_to_submit.append(hours)
        if "submitted" in first_at and "responded" in first_at:
            hours = _hours_between(first_at["submitted"], first_at["responded"])
            if hours is not None:
                submit_to_response.append(hours)
        if "submitted" in first_at:
            sub_at = _ts(first_at["submitted"])
            if sub_at and (datetime.now(timezone.utc) - sub_at) < timedelta(days=1):
                submitted_24h += 1
            if sub_at and (datetime.now(timezone.utc) - sub_at) < timedelta(days=7):
                submitted_7d += 1

    submitted_total = by_state.get("submitted", 0) + by_state.get("responded", 0) + by_state.get("rejected", 0)
    response_total = by_state.get("responded", 0) + by_state.get("rejected", 0)
    return {
        "total": sum(by_state.values()),
        "by_state": dict(by_state),
        "submitted_24h": submitted_24h,
        "submitted_7d": submitted_7d,
        "median_tailor_to_submit_hours": (
            round(statistics.median(tailor_to_submit), 2) if tailor_to_submit else None
        ),
        "median_submit_to_response_hours": (
            round(statistics.median(submit_to_response), 2) if submit_to_response else None
        ),
        "response_rate": (
            round(response_total / submitted_total, 3) if submitted_total else None
        ),
        "errors": error_reasons.most_common(10),
    }


def build_summary(excel_path: Path, state_path: Path) -> dict:
    return {
        "generated_at": datetime.now(timezone.utc).isoformat(timespec="seconds"),
        "scrape": scrape_metrics(excel_path),
        "applications": application_metrics(state_path),
    }
